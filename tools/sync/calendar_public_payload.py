#!/usr/bin/env python3
"""Master Calendario (xlsx) -> the public calendar copy read by the app (schema scd-calendar-public-v3).

Usage:
  python3 -I tools/sync/calendar_public_payload.py master.xlsx <source_modified_iso> > calendar.json
  python3 -I tools/sync/calendar_public_payload.py --self-test

Source: Drive file 1eXS-kY5PCHFg4CXiAErNAKtIVEMIKgWs ("CALENDARIO_PARTITE_SCD_ColicoDerviese_2026-27_UFFICIALE..."),
sheet "MASTER EVENTI". Target: the single public JSON copy, Drive file 1xSQ2MClw7LlIeWhJCSXllmCMQMrXYPn5, read by
apps/superapp/src/lib/calendar-bridge.server.ts and checked by calendar-mirror.ts (validateMirror): the app never reads
the master and never stores games in the database (NO_PARALLEL_CALENDAR_EDITING_IN_DB).

The rules are the ones of the hand-made copy of 09-10/10/2026 (projection SCD_RC6_GERARCHIA_FONTI_20261010), re-derived
on 15/10/2026 by comparing all 156 published games with the master row by row:
- out: the withdrawn U18 team, rows WITHDRAWN / REST / REJECTED, type RIPOSO, and club events (CLUB_EVENT) whose time
  or venue is not CONFIRMED (a provisional friendly is not published until the club confirms it in the master);
- fields: only the allowlisted keys of calendar-mirror.ts; verification flags from TIME_STATUS / VENUE_STATUS /
  CONFLICT_STATUS; notices from the same statuses; source FIGC/LND or SCD;
- one hand correction of 09-10/10 (official FIGC time over the master on 10/10, rule FIGC_LND_OFFICIAL_FIXTURES_WIN) is
  kept in OVERRIDES with its reason and the master value it corrects: it applies only while the master still carries
  that value (a newer master value wins and the correction is reported on stderr as stale);
- events are sorted by date, effective time and id.
A row whose category maps to no known team, an unreadable date or time, an over-long text or PII-like text stops the
run (exit 2): nothing is written. The checks mirror calendar-mirror.ts validateMirror, so the app never refuses a copy
this script produced.
No personal data: the master sheet has teams, venues and statuses only; the text checks below refuse e-mails, phones
and fiscal codes anyway (same patterns as the app validator).
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sys

SCHEMA = "scd-calendar-public-v3"
AUDIENCE = "PUBLIC_READ_ONLY_SNAPSHOT"
SOURCE_FILE_ID = "1eXS-kY5PCHFg4CXiAErNAKtIVEMIKgWs"
SEASON = "2026/27"
PROJECTION = "SCD_RC7_MASTER_SYNC"
SHEET = "MASTER EVENTI"

GROUPS: dict[str, dict] = {
    "2014": {"label": "Esordienti · 2014", "color": "#7654B8", "logo": None},
    "2015": {"label": "Esordienti · 2015", "color": "#148B99", "logo": "SCD_U12_2015.png"},
    "prima": {"label": "Prima Squadra", "color": "#AD7616", "logo": "SCD_PRIMA_SQUADRA.png"},
    "u19": {"label": "Juniores U19 Elite", "color": "#375B89", "logo": "SCD_JUNIORES_2008_2009.png"},
    "u16": {"label": "U16 Elite · 2011", "color": "#C85D22", "logo": "SCD_U16_2011.png"},
    "u15": {"label": "U15 · Giovanissimi", "color": "#C23C3F", "logo": "SCD_U15_2012_2013.png"},
    "pulcini": {"label": "Pulcini", "color": "#3269BC", "logo": None},
    "primicalci": {"label": "Primi Calci", "color": "#568C23", "logo": "SCD_PRIMI_2018_2019.png"},
    "piccoli": {"label": "Piccoli Amici", "color": "#2A9CB7", "logo": "SCD_PICCOLI_2020_2021.png"},
    "esordienti": {"label": "Esordienti · annata da verificare", "color": "#7E8799", "logo": None},
    "other": {"label": "Altre categorie", "color": "#668196", "logo": None},
}

# Category text of the master -> team group. Order matters (first match wins). None = never published.
CATEGORY_RULES: list[tuple[re.Pattern, str | None]] = [
    (re.compile(r"\bU\s*18\b", re.I), None),  # U18 withdrawn from the season (Master Calendario 2026/27)
    (re.compile(r"\bU\s*19\b|juniores", re.I), "u19"),
    (re.compile(r"prima\s+squadra", re.I), "prima"),
    (re.compile(r"\bU\s*16\b", re.I), "u16"),
    (re.compile(r"\bU\s*15\b|giovanissimi", re.I), "u15"),
    (re.compile(r"esordienti.*2015|2015.*esordienti", re.I), "2015"),
    (re.compile(r"esordienti.*2014|2014.*esordienti", re.I), "2014"),
    (re.compile(r"primi\s+calci", re.I), "primicalci"),
    (re.compile(r"pulcini", re.I), "pulcini"),
    (re.compile(r"piccoli\s+amici", re.I), "piccoli"),
]

EXCLUDED_STATUS = {"WITHDRAWN", "REST", "REJECTED"}
WITHDRAWN_RX = re.compile(r"RITIRAT|ANNULLAT|CANCEL|WITHDRAWN", re.I)  # same words the app validator refuses

# Hand corrections of the 09-10/10 copy (source hierarchy). "master" = the master value the correction overrides: the
# correction applies only while the master still says exactly that.
OVERRIDES: dict[str, dict] = {
    "EVT-U19E-5A-MEDA1913-FUORI": {
        "master": {"time": "17:00"},
        "set": {
            "time": "14:30",
            "end": "",
            "notice": [
                "Orario ufficiale FIGC/LND: 14:30, C.U. CRL n.28 del 24/09/2026, confermato nel programma gare CRL 09-12/10 pubblicato il 09/10/2026."
            ],
        },
    },
}

BAD_TEXT = [
    re.compile(r"</?[a-z!]", re.I),
    re.compile("[\u0000-\u0008\u000b-\u001f\u007f-\u009f\u200b-\u200f\u202a-\u202e\u2066-\u2069]"),
    re.compile(r"[^\s@]+@[^\s@]+\.[a-z]{2,}", re.I),
    re.compile(r"(\+39|0039)\s?\d"),
    re.compile(r"(?<!\d)3\d{2}[\s.-]?\d{6,7}(?!\d)"),
    re.compile(r"\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b", re.I),
    re.compile(r"\btarga\b|\bautista\b|\btelefono\b|\bcellulare\b|\biban\b", re.I),
    re.compile(r"docs\.google\.com/spreadsheets|" + SOURCE_FILE_ID),
]


class Stop(Exception):
    """A row the converter refuses to guess: the run stops and nothing is written."""


def text(v, limit: int = 500) -> str:
    if v is None:
        return ""
    s = re.sub(r"[ \t\r\n\u00a0]+", " ", str(v)).strip()
    if len(s) > limit:
        raise Stop(f"testo troppo lungo ({len(s)} > {limit}): {s[:40]}…")
    for r in BAD_TEXT:
        if r.search(s):
            raise Stop(f"testo non ammesso nel master: {s[:40]}…")
    return s


def iso_date(v) -> str:
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = text(v, 20)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m:
        return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        try:
            return dt.date.fromisoformat(s[:10]).isoformat()
        except ValueError:
            pass
    raise Stop(f"data non leggibile: {s!r}")


def hhmm(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, (dt.time, dt.datetime)):
        return v.strftime("%H:%M")
    s = text(v, 10)
    m = re.match(r"^(\d{1,2})[:.](\d{2})(:\d{2})?$", s)
    if not m or int(m[1]) > 23 or int(m[2]) > 59:
        raise Stop(f"ora non leggibile: {s!r}")
    return f"{int(m[1]):02d}:{m[2]}"


def group_of(category: str) -> str | None:
    for rx, g in CATEGORY_RULES:
        if rx.search(category):
            return g
    raise Stop(f"categoria senza squadra: {category!r} (aggiungerla in CATEGORY_RULES)")


def published(row: dict) -> bool:
    status = row.get("MATCH_STATUS") or ""
    if status in EXCLUDED_STATUS or WITHDRAWN_RX.search(str(status)):
        return False
    if (row.get("TYPE") or "").upper() == "RIPOSO":
        return False
    if row.get("MATCH_STATUS") == "CLUB_EVENT" and not (
        row.get("TIME_STATUS") == "CONFIRMED" and row.get("VENUE_STATUS") == "CONFIRMED"
    ):
        return False
    return True


def notices(row: dict) -> list[str]:
    out: list[str] = []
    if row.get("TIME_STATUS") != "CONFIRMED" or row.get("VENUE_STATUS") != "CONFIRMED":
        out.append("Orario o sede da verificare")
    if (row.get("CONFLICT_STATUS") or "NONE") != "NONE":
        out.append("Utilizzo struttura da verificare")
    elif row.get("OPERATIONAL_STATUS") == "OFFICIAL_TO_VERIFY":
        out.append("Attesa verifica")
    return out


def event_of(row: dict, quiet: bool = False) -> dict | None:
    category = text(row.get("CATEGORY"), 120)
    group = group_of(category)
    if group is None or not published(row):
        return None
    eid = text(row.get("EVENT_ID"), 120)
    ha = text(row.get("HOME_AWAY"), 10).upper()
    if ha not in ("CASA", "FUORI"):
        raise Stop(f"casa/fuori non leggibile per {eid}")
    e = {
        "id": eid,
        "date": iso_date(row.get("DATE")),
        "time": hhmm(row.get("TIME")),
        "end": hhmm(row.get("END_TIME")),
        "category": category,
        "group": group,
        "type": text(row.get("TYPE"), 40),
        "home_away": ha,
        "venue": text(row.get("VENUE")),
        "opponent": text(row.get("OPPONENT_EVENT")),
        "round": text(row.get("ROUND"), 40),
        "source": {"FIGC_LND": "FIGC/LND", "SCD_INTERNAL": "SCD"}.get(row.get("SOURCE_TYPE") or "", text(row.get("SOURCE_TYPE"), 40)),
        "notice": notices(row),
        "match_status": text(row.get("MATCH_STATUS"), 40),
        "is_variation": bool(text(row.get("VARIATION"))),
        "verification": {
            "time_confirmed": row.get("TIME_STATUS") == "CONFIRMED",
            "venue_confirmed": row.get("VENUE_STATUS") == "CONFIRMED",
            "facility_review_required": (row.get("CONFLICT_STATUS") or "NONE") != "NONE",
        },
    }
    fix = OVERRIDES.get(eid)
    if fix:
        if all(e.get(k) == v for k, v in fix["master"].items()):
            e.update(fix["set"])
        elif not quiet:
            print(f"AVVISO: correzione manuale di {eid} non applicata, il master è cambiato: vale il master", file=sys.stderr)
    return e


def build(rows: list[dict], source_modified_at: str, source_sha256: str) -> dict:
    events, seen = [], set()
    for r in rows:
        if not r.get("EVENT_ID"):
            continue
        e = event_of(r)
        if e is None:
            continue
        if e["id"] in seen:
            raise Stop(f"EVENT_ID duplicato: {e['id']}")
        seen.add(e["id"])
        events.append(e)
    if not events:
        raise Stop("nessuna gara pubblicabile: master vuoto o foglio rinominato")
    events.sort(key=lambda e: (e["date"], e["time"] or "99:99", e["id"]))
    return {
        "schema": SCHEMA,
        "source_modified_at": source_modified_at,
        "source_file_id": SOURCE_FILE_ID,
        "source_sha256": source_sha256,
        "season": SEASON,
        "snapshot_status": "COPIA_AUTOMATICA_DAL_MASTER",
        "groups": GROUPS,
        "events": events,
        "transport": {
            "confirmed_trips": 0,
            "status": "NESSUNA_CORSA_ASSEGNATA_NEL_MASTER_FLOTTA",
            "caption": "Iscrizioni pulmini e corse pianificate sono fonti separate.",
        },
        "facilities": {"reference": "Rotazione Campi Annuale, distinta dal Calendario", "allocations": "DA_VERIFICARE"},
        "projection_revision": PROJECTION,
        "audience": AUDIENCE,
    }


def read_master(path: str) -> list[dict]:
    import openpyxl  # only needed for real files; the self-test runs without it

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise Stop(f"foglio {SHEET!r} non trovato")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    head = [str(c).strip() if c is not None else "" for c in rows[0]]
    for need in ("EVENT_ID", "DATE", "TIME", "CATEGORY", "TYPE", "HOME_AWAY", "MATCH_STATUS", "TIME_STATUS", "VENUE_STATUS", "CONFLICT_STATUS"):
        if need not in head:
            raise Stop(f"colonna {need} mancante nel master")
    return [{h: r[i] if i < len(r) else None for i, h in enumerate(head) if h} for r in rows[1:] if r and any(c not in (None, "") for c in r)]


def self_test() -> None:
    base = {
        "EVENT_ID": "EVT-PRIMA-1A-X-CASA", "DATE": dt.datetime(2026, 10, 18), "TIME": dt.time(15, 30), "END_TIME": dt.time(17, 15),
        "CATEGORY": "Prima Squadra", "TYPE": "CAMPIONATO", "HOME_AWAY": "CASA", "VENUE": "COLICO - CAMPO 1", "OPPONENT_EVENT": "X",
        "ROUND": "1A", "SOURCE_TYPE": "FIGC_LND", "MATCH_STATUS": "OFFICIAL_CALENDAR", "TIME_STATUS": "CONFIRMED",
        "VENUE_STATUS": "CONFIRMED", "CONFLICT_STATUS": "NONE", "OPERATIONAL_STATUS": "READY", "VARIATION": None,
    }
    e = event_of(base)
    assert e and e["group"] == "prima" and e["time"] == "15:30" and e["end"] == "17:15" and e["notice"] == [] and e["source"] == "FIGC/LND"
    assert e["verification"] == {"time_confirmed": True, "venue_confirmed": True, "facility_review_required": False}
    assert event_of({**base, "CATEGORY": "U18 Élite"}) is None
    assert event_of({**base, "MATCH_STATUS": "WITHDRAWN"}) is None
    assert event_of({**base, "TYPE": "RIPOSO", "MATCH_STATUS": "REST"}) is None
    assert event_of({**base, "MATCH_STATUS": "CLUB_EVENT", "TIME_STATUS": "CLUB_PROVISIONAL", "VENUE_STATUS": "CLUB_PROVISIONAL"}) is None
    assert event_of({**base, "MATCH_STATUS": "CLUB_EVENT", "SOURCE_TYPE": "SCD_INTERNAL"})["source"] == "SCD"
    v = event_of({**base, "TIME_STATUS": "DEFAULT_TO_VERIFY", "VENUE_STATUS": "DEFAULT_TO_VERIFY", "OPERATIONAL_STATUS": "OFFICIAL_TO_VERIFY"})
    assert v["notice"] == ["Orario o sede da verificare", "Attesa verifica"] and v["verification"]["time_confirmed"] is False
    c = event_of({**base, "CONFLICT_STATUS": "HARD_CONFLICT", "OPERATIONAL_STATUS": "BLOCKED_CONFLICT"})
    assert c["notice"] == ["Utilizzo struttura da verificare"] and c["verification"]["facility_review_required"] is True
    assert event_of({**base, "CATEGORY": "Esordienti 12 anni (2014)"})["group"] == "2014"
    assert event_of({**base, "CATEGORY": "Pulcini misti 7v7"})["group"] == "pulcini"
    assert event_of({**base, "VENUE": "COLICO -\u00a0CAMPO 1"})["venue"] == "COLICO - CAMPO 1"
    o = event_of({**base, "EVENT_ID": "EVT-U19E-5A-MEDA1913-FUORI", "CATEGORY": "Juniores U19 Élite", "TIME": dt.time(17, 0)})
    assert o["time"] == "14:30" and o["end"] == "" and o["notice"][0].startswith("Orario ufficiale FIGC/LND")
    n = event_of({**base, "EVENT_ID": "EVT-U19E-5A-MEDA1913-FUORI", "CATEGORY": "Juniores U19 Élite", "TIME": dt.time(16, 0)}, quiet=True)
    assert n["time"] == "16:00" and n["notice"] == []  # newer master value wins over a stale correction
    for st in ("ANNULLATA", "CANCELLED", "RITIRATA"):
        assert event_of({**base, "MATCH_STATUS": st}) is None
    for bad in (
        {**base, "CATEGORY": "Calcio a 5"},
        {**base, "VENUE": "chiamare 333 1234567"},
        {**base, "HOME_AWAY": "?"},
        {**base, "TIME": "15,30"},
        {**base, "DATE": "2026-13-45"},
        {**base, "VENUE": "x" * 501},
        {**base, "OPPONENT_EVENT": "Squadra\u200b"},
    ):
        try:
            event_of(bad)
            raise AssertionError("doveva fermarsi")
        except Stop:
            pass
    try:
        build([{**base}, {**base}], "2026-10-15T09:47:09.544Z", "0" * 64)
        raise AssertionError("duplicato non fermato")
    except Stop:
        pass
    out = build([base, {**base, "EVENT_ID": "EVT-U18-Y", "CATEGORY": "U18 Élite"}], "2026-10-15T09:47:09.544Z", "0" * 64)
    assert len(out["events"]) == 1 and out["schema"] == SCHEMA and out["audience"] == AUDIENCE
    assert set(out["events"][0]) <= {"id", "date", "time", "end", "category", "group", "type", "home_away", "venue", "opponent", "round", "source", "notice", "match_status", "is_variation", "verification"}
    print("self-test ok")


def main(argv: list[str]) -> int:
    if argv[1:2] == ["--self-test"]:
        self_test()
        return 0
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 1
    path, modified = argv[1], argv[2]
    if not re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}", modified):
        print("BLOCKER: source_modified_iso non valido", file=sys.stderr)
        return 1
    try:
        with open(path, "rb") as f:
            sha = hashlib.sha256(f.read()).hexdigest()
        out = build(read_master(path), modified, sha)
    except Stop as e:
        print(f"BLOCKER: {e}", file=sys.stderr)
        return 2
    json.dump(out, sys.stdout, ensure_ascii=False, indent=1)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
