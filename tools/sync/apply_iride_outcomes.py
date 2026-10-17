#!/usr/bin/env python3
"""Iride v5 «conferme gare»: apply the outcomes registered by Iride to a copy of the Master Calendario (xlsx), surgically.

Usage:
  python3 -I tools/sync/apply_iride_outcomes.py master.xlsx outcomes.json out.xlsx [--today YYYY-MM-DD]
  python3 -I tools/sync/apply_iride_outcomes.py --fingerprint outcomes.json
  python3 -I tools/sync/apply_iride_outcomes.py --self-test

Where it runs (Francesco, 17/10/2026): the master stays an .xlsx on Drive (file 1eXS-kY5PCHFg4CXiAErNAKtIVEMIKgWs, sheet
"MASTER EVENTI") and is never written by automation. Iride (Make scenario 7979437, route CONFERME) records every
confirm / cancel / move of the Responsabile attività di base in the Make data store «SCD · Iride · Esiti gare
(registro)». The hourly task "SCD Calendario · copia pubblica oraria" (apps/superapp/docs/CALENDAR_HOURLY_TASK.md)
downloads the master, lists that data store, runs this script on the downloaded copy (overlay) and gives the patched copy
to calendar_public_payload.py, which publishes branch data/calendar-public. The outcomes stay in the data store and are
re-applied every hour; when the Segreteria writes the same change into the master, the outcome becomes a no-op or is
skipped ("orario cambiato nel master": the master always wins).

outcomes.json = the records of that data store as returned by Make (data-store-records_list):
[{"key": "...", "data": {"event_id", "azione", "data_gara", "ora_prima", "sede_prima", "data_nuova", "ora_nuova",
"sede_nuova", "categoria", "chi_ruolo", "quando", ...}}]. The key contains the WhatsApp message id (which encodes a phone
number): it is never printed or written, only sha256(key)[:10].

Why surgical: Make's Google Sheets modules cannot write an Office file, and saving it with openpyxl would drop shapes and
the cached values of the formulas. This script rewrites only the touched <c> cells of the MASTER EVENTI rows inside the
zip and copies every other part byte for byte; workbook.xml gets fullCalcOnLoad.

Mapping (Francesco 17/10: conferma = CONFIRMED, annulla = CANCELLED, sposta = RESCHEDULED with new date/time/pitch):
- CONFERMA: TIME_STATUS = CONFIRMED, VENUE_STATUS = CONFIRMED, OPERATIONAL_STATUS OFFICIAL_TO_VERIFY / CLUB_TO_VERIFY ->
            READY; MATCH_STATUS unchanged (OFFICIAL_CALENDAR / CLUB_EVENT keep telling families who set the game);
- ANNULLA:  MATCH_STATUS = CANCELLED (calendar_public_payload.py drops it from the app);
- SPOSTA:   MATCH_STATUS = RESCHEDULED; DATE / TIME / VENUE = the new values given; TIME_STATUS and VENUE_STATUS =
            CONFIRMED; VARIATION = "Variazione SCD dd/mm" (CAMBIA, the v5 draft name, is read as SPOSTA).
Always: NOTES += "Iride dd/mm: <azione> dal Resp. attività di base [IRIDE <hash>]", LAST_VERIFIED, ROW_VERSION + 1,
CHANGE_FLAG = IRIDE. TIME_STATUS / VENUE_STATUS become CONFIRMED only when the time / venue is not empty
(dd/mm = the day Iride registered it, so the hourly result is the same every hour).
Never touched: CONFLICT_STATUS (a field conflict stays open), other rows, other sheets.

Safety: an outcome is skipped (and reported) when the EVENT_ID is missing, the date, time or MATCH_STATUS in the master is
no longer the one Iride showed (match_status_prima when recorded; without it any cancelled / withdrawn status, since Iride
only lists published games), LAST_VERIFIED in the master is later than the day Iride registered the outcome (the
Segreteria worked on the row since: the master wins), the category is outside the attività di base
(up to annata 2014: Esordienti, Pulcini, Primi Calci, Piccoli Amici), the action / date / time / pitch is not valid
(new pitch only COLICO - CAMPO 1 / 2 and only for a home game; new date inside the season), or the [IRIDE <hash>] tag is
already in NOTES (idempotent). Exit 2 = nothing written.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

SHEET = "MASTER EVENTI"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
SCOPE_RX = re.compile(r"PRIMI\s+CALCI|PICCOLI\s+AMICI|PULCINI|ESORDIENTI", re.I)
HOME_VENUES = {"COLICO - CAMPO 1", "COLICO - CAMPO 2"}
TIME_RX = re.compile(r"^([01]\d|2[0-3]):([0-5]\d)$")
DATE_RX = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")
SEASON = (dt.date(2026, 7, 1), dt.date(2027, 6, 30))
ALIAS = {"CAMBIA": "SPOSTA"}
LABEL = {"CONFERMA": "confermata", "ANNULLA": "annullata", "SPOSTA": "spostata"}
GONE_RX = re.compile(r"RITIRAT|ANNULLAT|CANCEL|WITHDRAWN|^REST$|^REJECTED$", re.I)  # never in a list Iride sends
MAX_RECORDS = 100  # Make returns at most 100 records per list call: a full page may hide outcomes


class Stop(Exception):
    pass


def col_letters(ref: str) -> str:
    return re.match(r"[A-Z]+", ref).group(0)


def col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def sheet_path(z: zipfile.ZipFile) -> str:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rid = None
    for s in wb.find("m:sheets", NS):
        if s.get("name") == SHEET:
            rid = s.get("{%s}id" % NS["r"])
    if not rid:
        raise Stop(f"foglio {SHEET!r} non trovato")
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    for r in rels:
        if r.get("Id") == rid:
            t = r.get("Target")
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise Stop("relazione del foglio non trovata")


def shared_strings(z: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    root = ET.fromstring(z.read("xl/sharedStrings.xml"))
    return ["".join(t.text or "" for t in si.iter("{%s}t" % NS["m"])) for si in root.findall("m:si", NS)]


def cell_value(c: ET.Element, sst: list[str]):
    t = c.get("t")
    if t == "inlineStr":
        return "".join(x.text or "" for x in c.iter("{%s}t" % NS["m"]))
    v = c.find("m:v", NS)
    if v is None:
        return None
    if t == "s":
        return sst[int(v.text)]
    if t in ("str", "b", "e"):
        return v.text
    try:
        return float(v.text)
    except (TypeError, ValueError):
        return v.text


def as_date(v) -> str:
    if isinstance(v, float):
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).strftime("%d/%m")
    s = str(v or "").strip()
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"{int(m[1]):02d}/{int(m[2]):02d}"
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m[3]}/{m[2]}"
    return s


def as_time(v) -> str:
    if isinstance(v, float):
        mins = round((v % 1) * 24 * 60)
        return f"{mins // 60:02d}:{mins % 60:02d}"
    s = str(v or "").strip()
    m = re.match(r"^(\d{1,2})[:.](\d{2})", s)
    return f"{int(m[1]):02d}:{m[2]}" if m else s


def read_rows(xml: bytes, sst: list[str]):
    root = ET.fromstring(xml)
    data = root.find("m:sheetData", NS)
    rows = data.findall("m:row", NS)
    if not rows:
        raise Stop("foglio vuoto")
    head = {}
    for c in rows[0].findall("m:c", NS):
        v = cell_value(c, sst)
        if v not in (None, ""):
            head[str(v).strip()] = col_letters(c.get("r"))
    for need in ("EVENT_ID", "DATE", "TIME", "CATEGORY", "VENUE", "MATCH_STATUS", "TIME_STATUS", "VENUE_STATUS", "NOTES"):
        if need not in head:
            raise Stop(f"colonna {need} mancante nel master")
    by_id = {}
    for r in rows[1:]:
        vals = {col_letters(c.get("r")): (cell_value(c, sst), c.get("s")) for c in r.findall("m:c", NS)}
        rec = {h: vals.get(L, (None, None))[0] for h, L in head.items()}
        style = {h: vals.get(L, (None, None))[1] for h, L in head.items()}
        eid = str(rec.get("EVENT_ID") or "").strip()
        if eid:
            if eid in by_id:
                raise Stop(f"EVENT_ID duplicato nel master: {eid}")
            by_id[eid] = {"row": int(r.get("r")), "rec": rec, "style": style}
    return head, by_id


def new_cell(ref: str, style: str | None, value) -> str:
    s = f' s="{style}"' if style else ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{s}><v>{repr(float(value)) if isinstance(value, float) else value}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{esc(str(value))}</t></is></c>'


def set_cells(xml: str, row: int, cells: dict[str, str]) -> str:
    """cells: column letters -> full <c> xml. Replaces existing cells, inserts missing ones in column order."""
    rx_row = re.compile(r'(<row\b[^>]*\br="%d"[^>]*?)(/>|>(.*?)</row>)' % row, re.S)
    m = rx_row.search(xml)
    if not m:
        raise Stop(f"riga {row} non trovata nel foglio")
    inner = m.group(3) or ""
    parts = re.findall(r'<c\b[^>]*?(?:/>|>.*?</c>)', inner, re.S)
    if "".join(parts) != inner:
        raise Stop(f"riga {row}: struttura inattesa, nessuna modifica")
    current = {col_letters(re.search(r'\br="([A-Z]+\d+)"', p).group(1)): p for p in parts}
    current.update(cells)
    body = "".join(current[k] for k in sorted(current, key=col_index))
    open_tag = m.group(1)
    open_tag = re.sub(r'\s+spans="[^"]*"', "", open_tag)  # spans is optional; drop it rather than keep a wrong one
    return xml[: m.start()] + open_tag + ">" + body + "</row>" + xml[m.end():]


def key_tag(key: str) -> str:
    return "[IRIDE " + hashlib.sha256(key.encode()).hexdigest()[:10] + "]"


def fingerprint(outcomes: list[dict]) -> str:
    """Stable id of a set of outcomes (order-independent): the hourly task compares it with the one of the last commit."""
    canon = sorted(json.dumps({"k": hashlib.sha256(str(o.get("key", "")).encode()).hexdigest(), "d": o.get("data", {})},
                              sort_keys=True, ensure_ascii=False) for o in outcomes)
    return hashlib.sha256("\n".join(canon).encode()).hexdigest()[:16]


def day_of(d: dict, today: dt.date) -> dt.date:
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(d.get("quando") or ""))
    if m:
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            pass
    return today


def parse_new_date(s: str) -> dt.date | None:
    m = DATE_RX.match(s)
    if not m:
        return None
    try:
        d = dt.date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None
    return d if SEASON[0] <= d <= SEASON[1] else None


def as_day(v) -> dt.date | None:
    """LAST_VERIFIED of the master: Excel serial, dd/mm/yyyy or yyyy-mm-dd (with or without time)."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
        except (OverflowError, ValueError):
            return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    try:
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
        if m:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None
    return None


def plan(head, by_id, outcomes: list[dict], today: dt.date):
    applied, skipped, edits = [], [], {}
    for o in sorted(outcomes, key=lambda x: (str(x.get("data", {}).get("quando", "")), key_tag(str(x.get("key", ""))))):
        key, d = str(o.get("key", "")), o.get("data", {}) or {}
        tag = key_tag(key)
        eid = str(d.get("event_id") or "").strip()
        az = str(d.get("azione") or "").strip().upper()
        az = ALIAS.get(az, az)
        data_n = str(d.get("data_nuova") or "").strip()
        ora_n = str(d.get("ora_nuova") or "").strip()
        sede_n = str(d.get("sede_nuova") or "").strip().upper()
        why = None
        hit = by_id.get(eid)
        if not hit:
            why = "EVENT_ID non trovato"
        else:
            rec = hit["rec"]
            cur = edits.get(eid, {})
            notes = str(cur.get("NOTES", rec.get("NOTES") or ""))
            if tag in notes:
                why = "già applicato"
            elif az not in LABEL:
                why = f"azione non valida: {az!r}"
            elif not SCOPE_RX.search(str(rec.get("CATEGORY") or "")):
                why = "categoria fuori dall'attività di base"
            elif d.get("data_gara") and as_date(cur.get("_DATE", rec.get("DATE"))) != as_date(str(d["data_gara"]).strip()):
                why = "data cambiata nel master dopo la richiesta"
            elif d.get("ora_prima") and as_time(cur.get("_TIME", rec.get("TIME"))) != as_time(d["ora_prima"]):
                why = "orario cambiato nel master dopo la richiesta"
            elif d.get("match_status_prima") and str(cur.get("MATCH_STATUS", rec.get("MATCH_STATUS")) or "").strip().upper() \
                    != str(d["match_status_prima"]).strip().upper():
                why = "stato gara cambiato nel master dopo la richiesta"
            elif not d.get("match_status_prima") and GONE_RX.search(str(cur.get("MATCH_STATUS", rec.get("MATCH_STATUS")) or "").strip()):
                why = "stato gara cambiato nel master dopo la richiesta"
            elif as_day(rec.get("LAST_VERIFIED")) and as_day(rec.get("LAST_VERIFIED")) > day_of(d, dt.date.max):
                why = "riga aggiornata nel master dopo la richiesta"
            elif az == "SPOSTA":
                if not (data_n or ora_n or sede_n):
                    why = "spostamento senza data, orario né campo"
                elif data_n and parse_new_date(data_n) is None:
                    why = "data nuova non valida"
                elif ora_n and not TIME_RX.match(ora_n):
                    why = "orario non valido"
                elif sede_n and sede_n not in HOME_VENUES:
                    why = "campo non ammesso"
                elif sede_n and str(rec.get("HOME_AWAY") or "").strip().upper() != "CASA":
                    why = "cambio campo su una gara in trasferta"
        if why:
            skipped.append({"key_hash": tag, "event_id": eid, "motivo": why})
            continue
        e = edits.setdefault(eid, {"_row": hit["row"]})
        rec = hit["rec"]
        day = day_of(d, today)
        if az == "ANNULLA":
            e["MATCH_STATUS"] = "CANCELLED"
        else:
            # a missing time or venue is never «confirmed»: the public copy keeps «Orario o sede da verificare»
            if (ora_n if az == "SPOSTA" and ora_n else str(cur.get("_TIME", rec.get("TIME")) or "").strip()):
                e["TIME_STATUS"] = "CONFIRMED"
            if (sede_n if az == "SPOSTA" and sede_n else str(cur.get("VENUE", rec.get("VENUE")) or "").strip()):
                e["VENUE_STATUS"] = "CONFIRMED"
            if "OPERATIONAL_STATUS" in head and str(rec.get("OPERATIONAL_STATUS") or "") in ("OFFICIAL_TO_VERIFY", "CLUB_TO_VERIFY", ""):
                e["OPERATIONAL_STATUS"] = "READY"  # a field conflict (BLOCKED_CONFLICT) stays as it is
        if az == "SPOSTA":
            e["MATCH_STATUS"] = "RESCHEDULED"
            if data_n:
                nd = parse_new_date(data_n)
                e["DATE"] = ("date", nd)
                e["_DATE"] = nd.strftime("%d/%m/%Y")
            if ora_n:
                e["TIME"] = ("time", ora_n)
                e["_TIME"] = ora_n
            if sede_n:
                e["VENUE"] = sede_n
            if "VARIATION" in head:
                e["VARIATION"] = "Variazione SCD " + day.strftime("%d/%m")
        what = ", ".join(x for x in (data_n, ora_n, sede_n) if x)
        note = f"Iride {day.strftime('%d/%m')}: {LABEL[az]}" + (f" ({what})" if az == "SPOSTA" else "") \
            + " dal Resp. attività di base " + tag
        e["NOTES"] = (notes + " | " if notes else "") + note
        if "LAST_VERIFIED" in head:
            e["LAST_VERIFIED"] = day.strftime("%d/%m/%Y")
        if "ROW_VERSION" in head:
            try:
                rv = int(float(e.get("ROW_VERSION", rec.get("ROW_VERSION") or 0)))
            except (TypeError, ValueError):
                rv = 0
            e["ROW_VERSION"] = rv + 1
        if "CHANGE_FLAG" in head:
            e["CHANGE_FLAG"] = "IRIDE"
        applied.append({"key_hash": tag, "event_id": eid, "azione": az})
    return applied, skipped, edits


def typed(value, original):
    """DATE / TIME keep the representation of the original cell: a number (Excel serial, styled) or a text."""
    if isinstance(value, tuple):
        kind, v = value
        if kind == "date":
            return float((v - dt.date(1899, 12, 30)).days) if isinstance(original, float) else v.strftime("%d/%m/%Y")
        h, mi = TIME_RX.match(v).groups()
        # a number only if the master cell is already a (styled) number: a bare fraction in an empty cell would be
        # read back as 0.666… and stop the converter
        return (int(h) * 60 + int(mi)) / 1440 if isinstance(original, float) else v
    return value


def apply(master: str, outcomes: list[dict], out: str, today: dt.date) -> dict:
    with zipfile.ZipFile(master) as z:
        path = sheet_path(z)
        sst = shared_strings(z)
        xml_b = z.read(path)
        head, by_id = read_rows(xml_b, sst)
        applied, skipped, edits = plan(head, by_id, outcomes, today)
        xml = xml_b.decode("utf-8")
        for eid, e in edits.items():
            row, style, rec = e["_row"], by_id[eid]["style"], by_id[eid]["rec"]
            cells = {head[h]: new_cell(f"{head[h]}{row}", style.get(h), typed(v, rec.get(h)))
                     for h, v in e.items() if not h.startswith("_")}
            xml = set_cells(xml, row, cells)
        ET.fromstring(xml.encode("utf-8"))  # still well-formed, or Stop before writing
        wb = z.read("xl/workbook.xml").decode("utf-8")
        if edits:
            if "<calcPr" in wb:
                wb = re.sub(r"<calcPr\b([^>]*?)(/?)>", lambda m: "<calcPr" + re.sub(r'\s+fullCalcOnLoad="[^"]*"', "", m.group(1))
                            + ' fullCalcOnLoad="1"' + m.group(2) + ">", wb, count=1)
            else:
                wb = wb.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
        with zipfile.ZipFile(out, "w") as w:
            for info in z.infolist():
                data = z.read(info.filename)
                if info.filename == path:
                    data = xml.encode("utf-8")
                elif info.filename == "xl/workbook.xml":
                    data = wb.encode("utf-8")
                w.writestr(info, data, compress_type=info.compress_type)
    warn = []
    if len(outcomes) >= MAX_RECORDS:
        warn.append(f"registro esiti con {len(outcomes)} voci: la lista di Make si ferma a {MAX_RECORDS}, riportare gli esiti nel master e archiviarli")
    return {"applied": applied, "skipped": skipped, "sheet": SHEET, "fingerprint": fingerprint(outcomes), "avvisi": warn}


def load_outcomes(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        raw = json.load(f)
    if isinstance(raw, dict):  # tolerate {"records": [...]} or {"data": [...]}
        raw = raw.get("records", raw.get("data", raw))
    if not isinstance(raw, list) or not all(isinstance(o, dict) and isinstance(o.get("data", {}), dict) for o in raw):
        raise Stop("outcomes.json non è una lista di record {key, data}")
    return raw


def make_master(path: str, rows: list[list], head: list[str]) -> None:
    """Synthetic master with the real MASTER EVENTI columns (used by the self-test and the unit tests)."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "CALENDARIO SETTIMANALE"
    wb.active["A1"] = "vista"
    wb.active["B2"] = "='MASTER EVENTI'!D2"
    ws = wb.create_sheet(SHEET)
    ws.append(head)
    for r in rows:
        ws.append(r)
    wb.save(path)


HEAD = ["MATCH_ID", "EVENT_ID", "DATE", "TIME", "CATEGORY", "TYPE", "HOME_AWAY", "VENUE", "OPPONENT_EVENT", "ROUND",
        "SOURCE_ID", "SOURCE_TYPE", "SOURCE_STATUS", "MATCH_STATUS", "TIME_STATUS", "VENUE_STATUS", "VARIATION_WATCH",
        "CONFLICT_STATUS", "CONFLICT_GROUP_ID", "OPERATIONAL_STATUS", "SYNC_STATUS", "VARIATION", "NOTES", "SOURCE_ROW",
        "LAST_VERIFIED", "UPDATE_POLICY", "ROW_VERSION", "CHANGE_FLAG", "MANUAL_LOCK", "GOOGLE_EVENT_ID", "PARENT_EVENT_ID", "END_TIME"]


def master_row(eid, d, t, cat, ha, venue, status="OFFICIAL_CALENDAR", conflict="NONE"):
    r = {h: None for h in HEAD}
    r.update({"MATCH_ID": "M" + eid[3:], "EVENT_ID": eid, "DATE": d, "TIME": t, "CATEGORY": cat, "TYPE": "CAMPIONATO",
              "HOME_AWAY": ha, "VENUE": venue, "OPPONENT_EVENT": "AVV", "ROUND": "2A", "SOURCE_TYPE": "FIGC_LND",
              "MATCH_STATUS": status, "TIME_STATUS": "DEFAULT_TO_VERIFY", "VENUE_STATUS": "DEFAULT_TO_VERIFY",
              "CONFLICT_STATUS": conflict, "OPERATIONAL_STATUS": "OFFICIAL_TO_VERIFY", "NOTES": "nota originale", "ROW_VERSION": 1})
    return [r[h] for h in HEAD]


def outcome(key, eid, az, data, ora_prima, **kw):
    return {"key": key, "data": {"event_id": eid, "azione": az, "data_gara": data, "ora_prima": ora_prima,
                                 "quando": "2026-10-17T08:00:00.000Z", **kw}}


def self_test() -> None:
    import os
    import tempfile
    import openpyxl
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import calendar_public_payload as cpp

    tmp = tempfile.mkdtemp()
    src, out = os.path.join(tmp, "m.xlsx"), os.path.join(tmp, "o.xlsx")
    make_master(src, [
        master_row("EVT-PC78-2A-A-FUORI", dt.datetime(2026, 10, 17), dt.time(10, 30), "Primi Calci 7-8 anni", "FUORI", "Vercurago"),
        master_row("EVT-ES12-2A-B-CASA", dt.datetime(2026, 10, 17), dt.time(15, 0), "Esordienti 12 anni (2014)", "CASA", "COLICO - CAMPO 2",
                   conflict="HARD_CONFLICT"),
        master_row("EVT-PULMIX-2A-C-FUORI", dt.datetime(2026, 10, 17), dt.time(15, 15), "Pulcini Misti 7v7", "FUORI", "Rogeno"),
        master_row("EVT-MAN-20261018-001", dt.datetime(2026, 10, 18), dt.time(15, 0), "Pulcini Misti 7v7", "CASA", "COLICO - CAMPO 1",
                   status="CLUB_EVENT"),
        master_row("EVT-U16E-6A-D-CASA", dt.datetime(2026, 10, 17), dt.time(18, 0), "U16 Élite", "CASA", "COLICO - CAMPO 1"),
        master_row("EVT-PC78-3A-E-CASA", dt.datetime(2026, 10, 24), dt.time(16, 0), "Primi Calci 7-8 anni", "CASA", "COLICO - CAMPO 2"),
        master_row("EVT-PA-1A-F-CASA", dt.datetime(2026, 10, 25), dt.time(10, 0), "Piccoli Amici 5-6 anni", "CASA", "COLICO - CAMPO 1"),
    ], HEAD)
    outs = [
        outcome("wamid.A-1", "EVT-PC78-2A-A-FUORI", "CONFERMA", "17/10", "10:30"),
        outcome("wamid.A-3", "EVT-ES12-2A-B-CASA", "ANNULLA", "17/10", "15:00"),
        outcome("wamid.A-4", "EVT-PULMIX-2A-C-FUORI", "SPOSTA", "17/10", "15:15", ora_nuova="16:00"),
        outcome("wamid.A-5", "EVT-MAN-20261018-001", "CAMBIA", "18/10", "15:00", sede_nuova="COLICO - CAMPO 2"),
        outcome("wamid.A-10", "EVT-PA-1A-F-CASA", "SPOSTA", "25/10", "10:00", data_nuova="31/10/2026", ora_nuova="09:30"),
        outcome("wamid.A-6", "EVT-U16E-6A-D-CASA", "CONFERMA", "17/10", "18:00"),          # out of scope
        outcome("wamid.A-7", "EVT-PC78-3A-E-CASA", "CONFERMA", "24/10", "15:00"),          # master time changed since
        outcome("wamid.A-8", "EVT-NON-ESISTE", "CONFERMA", "17/10", "10:00"),
        outcome("wamid.A-9", "EVT-MAN-20261018-001", "SPOSTA", "18/10", "15:00", sede_nuova="DERVIO"),  # pitch not allowed
        outcome("wamid.A-11", "EVT-PULMIX-2A-C-FUORI", "SPOSTA", "17/10", "16:00", sede_nuova="COLICO - CAMPO 1"),  # away game
    ]
    rep = apply(src, outs, out, dt.date(2026, 10, 17))
    assert sorted(a["event_id"] for a in rep["applied"]) == sorted(["EVT-PC78-2A-A-FUORI", "EVT-ES12-2A-B-CASA", "EVT-PULMIX-2A-C-FUORI",
                                                                     "EVT-MAN-20261018-001", "EVT-PA-1A-F-CASA"]), rep
    why = {s["event_id"] + ":" + s["motivo"] for s in rep["skipped"]}
    assert why == {"EVT-U16E-6A-D-CASA:categoria fuori dall'attività di base",
                   "EVT-PC78-3A-E-CASA:orario cambiato nel master dopo la richiesta",
                   "EVT-NON-ESISTE:EVENT_ID non trovato",
                   "EVT-MAN-20261018-001:campo non ammesso",
                   "EVT-PULMIX-2A-C-FUORI:cambio campo su una gara in trasferta"}, why
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        p = sheet_path(a)
        for n in a.namelist():
            if n not in (p, "xl/workbook.xml"):
                assert a.read(n) == b.read(n), n
        assert b'fullCalcOnLoad="1"' in b.read("xl/workbook.xml")
        assert b"wamid" not in b.read(p)
    rows = cpp.read_master(out)
    by = {r["EVENT_ID"]: r for r in rows}
    ev = {r["EVENT_ID"]: cpp.event_of(r, quiet=True) for r in rows}
    assert ev["EVT-PC78-2A-A-FUORI"]["verification"]["time_confirmed"] is True and ev["EVT-PC78-2A-A-FUORI"]["notice"] == []
    assert ev["EVT-PC78-2A-A-FUORI"]["match_status"] == "OFFICIAL_CALENDAR"
    assert ev["EVT-ES12-2A-B-CASA"] is None and by["EVT-ES12-2A-B-CASA"]["MATCH_STATUS"] == "CANCELLED"
    assert ev["EVT-PULMIX-2A-C-FUORI"]["time"] == "16:00" and ev["EVT-PULMIX-2A-C-FUORI"]["is_variation"] is True
    assert ev["EVT-PULMIX-2A-C-FUORI"]["match_status"] == "RESCHEDULED" and ev["EVT-PULMIX-2A-C-FUORI"]["venue"] == "Rogeno"
    assert ev["EVT-MAN-20261018-001"]["venue"] == "COLICO - CAMPO 2" and ev["EVT-MAN-20261018-001"]["notice"] == []
    assert ev["EVT-PA-1A-F-CASA"]["date"] == "2026-10-31" and ev["EVT-PA-1A-F-CASA"]["time"] == "09:30"
    assert ev["EVT-U16E-6A-D-CASA"]["notice"] == ["Orario o sede da verificare", "Attesa verifica"]  # untouched
    assert by["EVT-PC78-2A-A-FUORI"]["ROW_VERSION"] == 2 and by["EVT-PC78-2A-A-FUORI"]["CHANGE_FLAG"] == "IRIDE"
    assert "nota originale | Iride 17/10: confermata dal Resp. attività di base [IRIDE " in by["EVT-PC78-2A-A-FUORI"]["NOTES"]
    out2 = os.path.join(tmp, "o2.xlsx")
    rep2 = apply(out, outs, out2, dt.date(2026, 10, 17))
    assert rep2["applied"] == [] and sum(1 for s in rep2["skipped"] if s["motivo"] == "già applicato") == 5, rep2
    wb2 = openpyxl.load_workbook(out)
    assert wb2["CALENDARIO SETTIMANALE"]["B2"].value == "='MASTER EVENTI'!D2"
    assert fingerprint(outs) == fingerprint(list(reversed(outs))) and rep["fingerprint"] == fingerprint(outs)
    print("self-test ok")


def main(argv: list[str]) -> int:
    if argv[1:2] == ["--self-test"]:
        self_test()
        return 0
    try:
        if argv[1:2] == ["--fingerprint"] and len(argv) == 3:
            print(fingerprint(load_outcomes(argv[2])))
            return 0
        if len(argv) < 4:
            print(__doc__, file=sys.stderr)
            return 2
        today = dt.date.today()
        if "--today" in argv:
            today = dt.date.fromisoformat(argv[argv.index("--today") + 1])
        rep = apply(argv[1], load_outcomes(argv[2]), argv[3], today)
    except Stop as exc:
        print(f"BLOCKER: {exc}", file=sys.stderr)
        return 2
    for w in rep["avvisi"]:
        print(f"AVVISO: {w}", file=sys.stderr)
    print(json.dumps(rep, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
