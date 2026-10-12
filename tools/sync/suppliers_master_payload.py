#!/usr/bin/env python3
"""TESORERIA FORNITORI (xlsx export) -> JSON payload for public.scd_supplier_sync_copy (r80, machine M11).

Usage: python3 -I tools/sync/suppliers_master_payload.py tesoreria_fornitori.xlsx [source_modified_iso] > payload.json
       python3 -I tools/sync/suppliers_master_payload.py --self-test

Source (decision N2): Google Sheet "SCD - TESORERIA FORNITORI | SCADENZE E CHIUSURE 2026-2027"
(1l9_mmmo9zuALF2IrQvevhBFaNYFCWzOOjjyDuphQGsY), tabs FORNITORI (one row per supplier) and PAGAMENTI (agenda of
accepted instalments). Privacy: the payload never carries the CONTATTO and NOTE columns; e-mail addresses, phone
numbers and IBAN found in free text are removed. Supplier names are companies (organisation-level data).
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
import unicodedata

import openpyxl

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
IBAN = re.compile(r"\bIT\s?\d{2}\s?[A-Z](\s?[0-9A-Z]){22}\b", re.IGNORECASE)
PHONE = re.compile(r"(\+?\d[\d\s./-]{7,}\d)")
GMAIL = re.compile(r"^https://mail\.google\.com/")


def text(v, limit: int = 200) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s[:limit] or None


def clean(v, limit: int = 200) -> str | None:
    s = text(v, 400)
    if not s:
        return None
    s = EMAIL.sub("[mail]", s)
    s = IBAN.sub("[iban]", s)
    s = PHONE.sub(lambda m: "[tel]" if len(re.sub(r"\D", "", m[0])) >= 9 else m[0], s)
    return s[:limit].strip() or None


def money(v) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = (text(v, 30) or "").replace("€", "").replace(" ", "")
    if not s or s in ("-", "—"):
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def whole(v) -> int | None:
    m = money(v)
    return int(m) if m is not None and m >= 0 and float(m).is_integer() else None


def iso_date(v) -> str | None:
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = text(v, 20)
    if not s:
        return None
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if m:
        y = int(m[3]) + (2000 if len(m[3]) == 2 else 0)
        try:
            return dt.date(y, int(m[2]), int(m[1])).isoformat()
        except ValueError:
            return None
    return s[:10] if re.match(r"^\d{4}-\d{2}-\d{2}$", s[:10]) else None


def slug(name: str) -> str:
    n = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", n).strip("-")[:60]


def plan_group(status: str | None, residual: float | None) -> str:
    s = (status or "").upper()
    if "CONTEST" in s:
        return "CONTESTATO"
    if "CHIUS" in s or "SALDAT" in s or (residual is not None and residual <= 0 and "ACCETT" not in s):
        return "CHIUSO"
    if "ACCETT" in s or "ATTIV" in s or "IN CORSO" in s:
        return "PIANO_ATTIVO"
    return "DA_DEFINIRE"


def link_of(cell) -> str | None:
    """GMAIL column: a URL as value, or a hyperlink on the cell text."""
    v = text(cell.value, 500)
    if v and GMAIL.match(v):
        return v
    target = getattr(getattr(cell, "hyperlink", None), "target", None)
    return target if target and GMAIL.match(target) else None


def sheet_rows(ws) -> tuple[list[str], list[tuple], str | None]:
    """Header row = first row that contains FORNITORE among its first cells. Also returns the 'Aggiornato al' date."""
    rows = list(ws.iter_rows())
    aligned = None
    for r in rows[:6]:
        vals = [c.value for c in r]
        if vals and isinstance(vals[0], str) and vals[0].strip().lower().startswith("aggiornato al"):
            aligned = iso_date(vals[1] if len(vals) > 1 else None)
    hi = next(i for i, r in enumerate(rows[:12]) if any(isinstance(c.value, str) and c.value.strip().upper() == "FORNITORE" for c in r))
    head = [text(c.value, 60) or "" for c in rows[hi]]
    head = [h.upper() for h in head]
    body = [tuple(r) for r in rows[hi + 1:] if any(c.value not in (None, "") for c in r)]
    return head, body, aligned


def col(head: list[str], row: tuple, name: str):
    try:
        i = head.index(name)
    except ValueError:
        return None
    return row[i] if i < len(row) else None


def build(path: str, source_modified: str | None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    head, body, aligned = sheet_rows(wb["FORNITORI"])
    suppliers = []
    for r in body:
        name = text(getattr(col(head, r, "FORNITORE"), "value", None), 120)
        if not name:
            continue
        v = lambda h: getattr(col(head, r, h), "value", None)
        status = text(v("STATO PIANO"), 80)
        residual = money(v("RESIDUO"))
        gm = col(head, r, "GMAIL")
        suppliers.append({
            "key": slug(name),
            "supplier": name,
            "planStatus": status,
            "planGroup": plan_group(status, residual),
            "balance": money(v("SALDO")),
            "paid": money(v("PAGATO")),
            "residual": residual,
            "instalments": whole(v("N. RATE")),
            "nextAmount": money(v("PROSSIMA RATA")),
            "nextDue": iso_date(v("SCADENZA")),
            "closingExpected": iso_date(v("CHIUSURA PREVISTA")),
            "nextAction": clean(v("PROSSIMA AZIONE")),
            "commercialPosition": clean(v("POSIZIONE COMMERCIALE"), 120),
            "strategy": clean(v("STRATEGIA")),
            "mailUrl": link_of(gm) if gm is not None else None,
        })

    phead, pbody, paligned = sheet_rows(wb["PAGAMENTI"])
    instalments = []
    seen: dict[str, int] = {}
    for r in pbody:
        v = lambda h: getattr(col(phead, r, h), "value", None)
        name = text(v("FORNITORE"), 120)
        if not name:
            continue
        label = text(v("RATA"), 40)
        due = iso_date(v("SCADENZA"))
        base = f"{slug(name)}:{slug(label or '')}:{due or ''}"
        seen[base] = seen.get(base, 0) + 1
        gm = col(phead, r, "GMAIL")
        instalments.append({
            "key": base if seen[base] == 1 else f"{base}:{seen[base]}",
            "supplierKey": slug(name),
            "supplier": name,
            "label": label,
            "amount": money(v("IMPORTO")),
            "due": due,
            "planStatus": text(v("STATO PIANO"), 80),
            "urgency": text(v("URGENZA"), 40),
            "todo": clean(v("COSA FARE")),
            "mailUrl": link_of(gm) if gm is not None else None,
        })
    return {
        "source": "SCD - TESORERIA FORNITORI | SCADENZE E CHIUSURE 2026-2027",
        "season": "2026/27",
        "sourceModifiedAt": source_modified,
        "alignedOn": aligned or paligned,
        "suppliers": suppliers,
        "instalments": instalments,
    }


def self_test() -> None:
    import os
    import tempfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FORNITORI"
    ws.append(["FORNITORI - QUADRO UNICO"])
    ws.append(["Una riga per fornitore."])
    ws.append(["Aggiornato al", "30/09/2026"])
    ws.append(["FORNITORE", "STATO PIANO", "SALDO", "PAGATO", "RESIDUO", "N. RATE", "PROSSIMA RATA", "SCADENZA",
               "CHIUSURA PREVISTA", "CONTATTO", "PROSSIMA AZIONE", "POSIZIONE COMMERCIALE", "STRATEGIA", "GMAIL", "NOTE"])
    ws.append(["Fornitore Alfa Srl", "PIANO ACCETTATO", "1.000,00", 250, 750, 4, 250, dt.datetime(2026, 10, 31),
               "31/12/2026", "Mario Rossi 333 1234567", "Bonifico rata 2 entro 31/10/2026, scrivere a amm@alfa.it", "Attivo", "Rientro", "x", "nota privata"])
    ws["N5"].hyperlink = "https://mail.google.com/mail/u/0/#all/abc"
    ws.append(["Fornitore Beta", "CHIUSO", 300, 300, 0, 1, None, None, None, None, None, None, None, None, None])
    ws.append([None] * 15)
    p = wb.create_sheet("PAGAMENTI")
    p.append(["AGENDA PAGAMENTI FORNITORI"])
    p.append(["Solo piani accettati."])
    p.append(["Aggiornato al", "30/09/2026"])
    p.append(["URGENZA", "FORNITORE", "RATA", "IMPORTO", "SCADENZA", "PROMEMORIA D-5", "STATO PIANO", "RESIDUO FORNITORE",
              "COSA FARE", "GMAIL", "CONTATTO", "NOTE"])
    p.append(["ENTRO 5 GG", "Fornitore Alfa Srl", "2/4", 250, "31/10/2026", "26/10/2026", "PIANO ACCETTATO", 750,
              "Pagare su IT60X0542811101000000123456", "https://mail.google.com/mail/u/0/#all/abc", "x", "y"])
    p.append(["", "Fornitore Alfa Srl", "2/4", 250, "31/10/2026", None, "PIANO ACCETTATO", 750, None, None, None, None])
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(path)
        out = build(path, "2026-10-12T10:00:00Z")
    finally:
        os.unlink(path)
    s = out["suppliers"]
    assert out["alignedOn"] == "2026-09-30", out["alignedOn"]
    assert [x["key"] for x in s] == ["fornitore-alfa-srl", "fornitore-beta"]
    a = s[0]
    assert a["balance"] == 1000.0 and a["residual"] == 750 and a["instalments"] == 4
    assert a["nextDue"] == "2026-10-31" and a["closingExpected"] == "2026-12-31"
    assert a["planGroup"] == "PIANO_ATTIVO" and s[1]["planGroup"] == "CHIUSO"
    assert a["mailUrl"] == "https://mail.google.com/mail/u/0/#all/abc"
    assert a["nextAction"] == "Bonifico rata 2 entro 31/10/2026, scrivere a [mail]", a["nextAction"]
    assert "amm@alfa.it" not in json.dumps(out) and "Mario" not in json.dumps(out) and "nota privata" not in json.dumps(out)
    i = out["instalments"]
    assert len(i) == 2 and i[0]["key"] != i[1]["key"]
    assert "IT60" not in (i[0]["todo"] or "") and "[iban]" in (i[0]["todo"] or "")
    assert i[0]["due"] == "2026-10-31" and i[0]["amount"] == 250
    print("self-test ok")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--self-test":
        self_test()
        sys.exit(0)
    if len(sys.argv) < 2:
        sys.exit("usage: suppliers_master_payload.py file.xlsx [source_modified_iso] | --self-test")
    print(json.dumps(build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None), ensure_ascii=False))
